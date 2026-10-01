"""Mega-resume ingestion: resources/resume/* -> data/profile.yaml.

The LLM only transcribes and structures. Ids are assigned here, deterministically, and a
token-level post-check flags anything in the structured output that isn't in the source text.
On re-ingest, user-curated fields (tags/context/strength) and ids of matching items survive,
and a unified diff is produced for review.
"""

from __future__ import annotations

import difflib
import hashlib
import json
import os
import re
import shutil
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import yaml
from pydantic import BaseModel
from rapidfuzz import fuzz, process

from recrute.paths import Paths
from recrute.schemas import (
    Certification,
    Education,
    Experience,
    Profile,
    ProfileItem,
    Project,
    VerifierFlag,
)
from recrute.tailor.common import (
    STR,
    STRS,
    Completer,
    SupportIndex,
    arr,
    find_unsupported,
    obj,
    parse_llm,
    slugify,
)

RESUME_SUFFIXES = {".md", ".markdown", ".txt", ".docx", ".pdf"}


def profile_path(paths: Paths) -> Path:
    return paths.data / "profile.yaml"


def proposed_profile_path(paths: Paths) -> Path:
    return paths.data / "profile.proposed.yaml"


# --------------------------------------------------------------------------- reading sources


def extract_text(path: Path) -> str:
    """Plain text of a .md/.txt/.docx/.pdf file."""
    suffix = path.suffix.lower()
    if suffix == ".docx":
        import docx

        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    if suffix == ".pdf":
        from pypdf import PdfReader

        return "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
    return path.read_text(encoding="utf-8", errors="replace")


def read_sources(directory: Path) -> list[tuple[str, str]]:
    """[(file name, text)] for every supported file in `directory` (recursively, sorted)."""
    if not directory.is_dir():
        return []
    out = []
    for p in sorted(directory.rglob("*")):
        if p.is_file() and p.suffix.lower() in RESUME_SUFFIXES and not p.name.startswith("."):
            text = extract_text(p).strip()
            if text:
                out.append((p.relative_to(directory).as_posix(), text))
    return out


# --------------------------------------------------------------------------- LLM extraction

_ITEM = obj(text=STR, tags=STRS, metrics=STRS, context=STR)
EXTRACT_SCHEMA = obj(
    name=STR, headline=STR, email=STR, phone=STR, location=STR,
    links=arr(obj(label=STR, url=STR)),
    summary=STR,
    skills=arr(obj(category=STR, items=STRS)),
    experience=arr(obj(company=STR, title=STR, location=STR, start=STR, end=STR, summary=STR,
                       bullets=arr(_ITEM))),
    projects=arr(obj(name=STR, url=STR, summary=STR, tech=STRS, bullets=arr(_ITEM))),
    education=arr(obj(school=STR, degree=STR, field=STR, start=STR, end=STR, gpa=STR,
                      details=STRS)),
    certifications=arr(obj(name=STR, issuer=STR, date=STR, credential_id=STR)),
    awards=arr(_ITEM),
    extra=arr(_ITEM),
)

EXTRACT_SYSTEM = (
    "You transcribe resume documents into structured JSON. You are a transcriber, not a writer: "
    "every value must come from the source text. Never add, infer, embellish, merge or drop "
    "facts. Keep bullet wording verbatim and numbers exactly as written. Use \"\" or [] when the "
    "source lacks a field."
)

EXTRACT_PROMPT = """Structure the resume source below.
- experience: one entry per role (same company + different title = separate entries), most
  recent first. start/end as written ("present" if current). bullets verbatim.
- item.metrics: numeric results quoted from that bullet (e.g. "38%", "$2M"); else [].
- item.tags: 1-4 short lowercase topic labels (e.g. detection, cloud, ml, leadership).
- item.context: extra source detail about that bullet (notes, backstory); else "".
- projects: name = short project name only (tagline -> summary); tech = tools listed for
  the project. skills grouped by the source's categories
  (use "General" if none). links: label = linkedin|github|portfolio|other.
- awards: honors/competitions. extra: talks, publications, volunteering, other items.
- Several files may describe the same thing: merge duplicates, keep every distinct fact.

SOURCE:
{source}"""


class _Item(BaseModel):
    text: str
    tags: list[str] = []
    metrics: list[str] = []
    context: str = ""


class _Exp(BaseModel):
    company: str
    title: str
    location: str = ""
    start: str = ""
    end: str = ""
    summary: str = ""
    bullets: list[_Item] = []


class _Proj(BaseModel):
    name: str
    url: str = ""
    summary: str = ""
    tech: list[str] = []
    bullets: list[_Item] = []


class _Link(BaseModel):
    label: str
    url: str


class _Skills(BaseModel):
    category: str
    items: list[str]


class _Edu(BaseModel):
    school: str
    degree: str = ""
    field: str = ""
    start: str = ""
    end: str = ""
    gpa: str = ""
    details: list[str] = []


class _Cert(BaseModel):
    name: str
    issuer: str = ""
    date: str = ""
    credential_id: str = ""


class _Extracted(BaseModel):
    name: str = ""
    headline: str = ""
    email: str = ""
    phone: str = ""
    location: str = ""
    links: list[_Link] = []
    summary: str = ""
    skills: list[_Skills] = []
    experience: list[_Exp] = []
    projects: list[_Proj] = []
    education: list[_Edu] = []
    certifications: list[_Cert] = []
    awards: list[_Item] = []
    extra: list[_Item] = []


def _unique(base: str, used: set[str]) -> str:
    cand, n = base, 2
    while cand in used:
        cand, n = f"{base}-{n}", n + 1
    used.add(cand)
    return cand


def _items(prefix: str, raw: list[_Item], used: set[str]) -> list[ProfileItem]:
    out = []
    for i, it in enumerate(raw, 1):
        if not it.text.strip():
            continue
        out.append(ProfileItem(id=_unique(f"{prefix}-b{i}", used), text=it.text.strip(),
                               tags=[t.strip().lower() for t in it.tags if t.strip()],
                               metrics=[m for m in it.metrics if m.strip()],
                               context=it.context.strip()))
    return out


def to_profile(ex: _Extracted) -> Profile:
    """Convert the LLM's transcription into a Profile with deterministic ids."""
    used: set[str] = set()
    experience = []
    for e in ex.experience:
        eid = f"exp-{slugify(e.company, 40)}"
        if eid in used:
            eid = f"{eid}-{slugify(e.title, 40)}"
        eid = _unique(eid, used)
        experience.append(Experience(id=eid, company=e.company, title=e.title,
                                     location=e.location, start=e.start, end=e.end,
                                     summary=e.summary, bullets=_items(eid, e.bullets, used)))
    projects = []
    for p in ex.projects:
        short = re.split(r"\s+[—–:|-]\s+|\s*\(", p.name)[0] or p.name
        pid = _unique(f"proj-{slugify(short, 40)}", used)
        projects.append(Project(id=pid, name=p.name, url=p.url, summary=p.summary, tech=p.tech,
                                bullets=_items(pid, p.bullets, used)))
    education = [
        Education(id=_unique(f"edu-{slugify(ed.school, 40)}", used), **ed.model_dump())
        for ed in ex.education
    ]
    certs = [
        Certification(id=_unique(f"cert-{slugify(c.name, 40)}", used), **c.model_dump())
        for c in ex.certifications
    ]

    def standalone(prefix: str, raw: list[_Item]) -> list[ProfileItem]:
        return [
            ProfileItem(id=_unique(f"{prefix}-{slugify(it.text, 24)}", used), text=it.text.strip(),
                        tags=[t.lower() for t in it.tags], metrics=it.metrics,
                        context=it.context.strip())
            for it in raw if it.text.strip()
        ]

    skills: dict[str, list[str]] = {}
    for group in ex.skills:
        items = [s.strip() for s in group.items if s.strip()]
        if items:
            skills.setdefault(group.category.strip() or "General", []).extend(items)
    links: dict[str, str] = {}
    for link in ex.links:
        if link.url.strip():
            links[_unique(slugify(link.label or "other", 20), set(links))] = link.url.strip()
    return Profile(
        name=ex.name, headline=ex.headline, email=ex.email, phone=ex.phone, location=ex.location,
        links=links, summary=ex.summary, skills=skills, experience=experience,
        projects=projects, education=education, certifications=certs,
        awards=standalone("award", ex.awards), extra=standalone("extra", ex.extra),
    )



# --------------------------------------------------------------------------- post-check

_MONTHS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}
_MONTH_YEAR_RE = re.compile(r"\b(jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\.?,?"
                            r"\s*'?(\d{4})\b", re.IGNORECASE)
_YEAR_MONTH_RE = re.compile(r"\b((?:19|20)\d{2})[-/.](\d{1,2})\b")
_MONTH_NUM_YEAR_RE = re.compile(r"\b(\d{1,2})[-/.]((?:19|20)\d{2})\b")
_YEAR_RE = re.compile(r"\b((?:19|20)\d{2})\b")
_OPEN_ENDED = {"", "present", "current", "now", "ongoing", "today"}
_PHONE_RE = re.compile(r"\+?\d[\d\s().\-]{6,}\d")
_NAME_STOP = {"and", "of", "the", "at", "in", "for", "a", "an", "on", "to", "with"}


def date_keys(text: str) -> set[tuple[int, int | None]]:
    """{(year, month|None)} for every date in `text` ("May 2024", "2024-05", "05/2024", "2024")."""
    keys: set[tuple[int, int | None]] = set()
    for m in _MONTH_YEAR_RE.finditer(text):
        keys.add((int(m.group(2)), _MONTHS[m.group(1).lower()[:3]]))
    for m in _YEAR_MONTH_RE.finditer(text):
        if 1 <= int(m.group(2)) <= 12:
            keys.add((int(m.group(1)), int(m.group(2))))
    for m in _MONTH_NUM_YEAR_RE.finditer(text):
        if 1 <= int(m.group(1)) <= 12:
            keys.add((int(m.group(2)), int(m.group(1))))
    keys |= {(int(y), None) for y in _YEAR_RE.findall(text)}
    return keys


def _norm_url(url: str) -> str:
    u = re.sub(r"^[a-z][a-z0-9+.-]*://", "", url.strip().lower())
    return u.removeprefix("www.").rstrip("/")


def _missing_words(text: str, index: SupportIndex) -> list[str]:
    words = re.findall(r"[A-Za-z][A-Za-z0-9'+#&-]*", text)
    return [w for w in words if len(w) > 1 and w.lower() not in _NAME_STOP
            and not index.has_term(w.strip("'-"))]


def _source_window(text: str, lines: list[str], windows: list[str]) -> str | None:
    """The source lines around where `text` was most likely transcribed from."""
    probe = " ".join(text.split())[:120]
    best = process.extractOne(probe, windows, scorer=fuzz.partial_ratio, score_cutoff=70)
    if best is None:
        return None
    i = best[2]
    return "\n".join(lines[max(0, i - 1): i + 6])


def check_against_source(profile: Profile, source_text: str, *,
                         trusted: frozenset[str] | set[str] = frozenset()) -> list[VerifierFlag]:
    """Flag anything in the structured profile that isn't in the source text.

    - identity: name words, email, phone digits and link URLs (normalized) must occur;
    - names/places: company, title, school, field, certification, location words must occur;
    - dates: every (year, month) of employment/education/certification dates must occur;
    - free text: numbers and tech terms in bullets are checked against the source lines they
      were transcribed from, so a number found elsewhere in the resume is still flagged.
    `trusted`: strings the user curated earlier (preserved `context`), which aren't in the
    source by design.
    """
    index = SupportIndex.of([source_text])
    lower = source_text.lower()
    url_text = re.sub(r"[a-z][a-z0-9+.-]*://(www\.)?|\bwww\.", "", lower)
    src_dates = date_keys(source_text)
    src_phones = {re.sub(r"\D", "", m)[-10:] for m in _PHONE_RE.findall(source_text)}
    lines = [ln.strip() for ln in source_text.splitlines() if ln.strip()]
    windows = [" ".join(lines[i:i + 3]) for i in range(len(lines))]
    flags: list[VerifierFlag] = []

    def block(where: str, text: str, reason: str) -> None:
        flags.append(VerifierFlag(where=where, text=text, reason=reason, severity="block"))

    def check(where: str, text: str, local: str | None = None) -> None:
        near = SupportIndex.of([local]) if local else index
        misses = find_unsupported(text, near)
        if not misses:
            return
        anywhere = {u.token for u in find_unsupported(text, index)} if local else None
        for u in misses:
            if anywhere is not None and u.token not in anywhere:
                flags.append(VerifierFlag(
                    where=where, text=u.token, severity="warn",
                    reason=f"'{u.token}' is in the resume source, but not near this item"))
            else:
                block(where, u.token, f"'{u.token}' does not appear in the resume source")

    def words(where: str, text: str) -> None:
        for w in _missing_words(text, index):
            block(where, w, f"'{w}' does not appear in the resume source")

    def dates(where: str, value: str, header: str) -> None:
        """Dates must appear near the entry they belong to (its header's source lines)."""
        if value.strip().lower() in _OPEN_ENDED:
            return
        keys = date_keys(value)
        if not keys:
            check(where, value)
            return
        local = _source_window(header, lines, windows)
        if not keys <= src_dates:
            block(where, value, f"date '{value}' does not appear in the resume source")
        elif local is not None and not keys <= date_keys(local):
            flags.append(VerifierFlag(
                where=where, text=value, severity="block",
                reason=f"date '{value}' is in the resume source, but not near this entry"))

    # identity
    words("profile.name", profile.name)
    if profile.email and profile.email.strip().lower() not in lower:
        block("profile.email", profile.email, "email address not in the resume source")
    if profile.phone:
        digits = re.sub(r"\D", "", profile.phone)[-10:]
        if digits and digits not in src_phones:
            block("profile.phone", profile.phone, "phone number not in the resume source")
    for label, url in profile.links.items():
        if _norm_url(url) not in url_text:
            block(f"profile.links:{label}", url, "link not in the resume source")
    words("profile.location", profile.location)
    check("profile.summary", profile.summary)
    check("profile.headline", profile.headline)

    for item_id, item in profile.all_items().items():
        extra = [] if item.context in trusted else [item.context]
        check(f"profile:{item_id}", " ".join([item.text, *item.metrics, *extra]),
              _source_window(item.text, lines, windows))
    for e in profile.experience:
        words(f"profile:{e.id}", f"{e.company} {e.title} {e.location}")
        check(f"profile:{e.id}", e.summary)
        dates(f"profile:{e.id}.start", e.start, f"{e.title} {e.company}")
        dates(f"profile:{e.id}.end", e.end, f"{e.title} {e.company}")
    for p in profile.projects:
        words(f"profile:{p.id}", p.name)
        check(f"profile:{p.id}", " ".join([p.summary, *p.tech]))
        if p.url and _norm_url(p.url) not in url_text:
            block(f"profile:{p.id}.url", p.url, "link not in the resume source")
    for ed in profile.education:
        words(f"profile:{ed.id}", f"{ed.school} {ed.field}")
        check(f"profile:{ed.id}", " ".join([ed.degree, ed.gpa, *ed.details]))
        dates(f"profile:{ed.id}.start", ed.start, f"{ed.degree} {ed.field} {ed.school}")
        dates(f"profile:{ed.id}.end", ed.end, f"{ed.degree} {ed.field} {ed.school}")
    for c in profile.certifications:
        words(f"profile:{c.id}", f"{c.name} {c.issuer}")
        check(f"profile:{c.id}", c.credential_id)
        dates(f"profile:{c.id}.date", c.date, c.name)
    for cat, skills in profile.skills.items():
        check(f"profile.skills:{cat}", ", ".join(skills))
    return flags


# --------------------------------------------------------------------------- merge with existing


def _similar(a: str, b: str) -> float:
    return fuzz.token_sort_ratio(a.lower(), b.lower())


def _same_date(a: str, b: str) -> bool:
    ka, kb = date_keys(a), date_keys(b)
    if ka or kb:
        return ka == kb
    return a.strip().lower() == b.strip().lower()


class _Ids:
    """Id allocation for a merge. `reserved` are the old profile's ids: only the new item matched
    to that old item may take one, so a removed item's id is never recycled for new content."""

    def __init__(self, reserved: list[str]):
        self.reserved = set(reserved)
        self.taken: set[str] = set()

    def free(self, candidate: str) -> bool:
        return candidate not in self.taken and candidate not in self.reserved

    def adopt(self, old_id: str) -> str:
        self.taken.add(old_id)
        return old_id

    def fresh(self, base: str) -> str:
        cand, n = base, 2
        while not self.free(cand):
            cand, n = f"{base}-{n}", n + 1
        self.taken.add(cand)
        return cand

    def next_bullet(self, prefix: str) -> str:
        n = 1
        while not self.free(f"{prefix}-b{n}"):
            n += 1
        return self.adopt(f"{prefix}-b{n}")


def _pair(new: list[Any], old: list[Any], score: Any) -> dict[int, Any]:
    """One-to-one matching (best scores first): new index -> old entry."""
    cands = sorted(((s, i, j) for i, n in enumerate(new) for j, o in enumerate(old)
                    if (s := score(n, o)) is not None), key=lambda c: -c[0])
    pairs: dict[int, Any] = {}
    used_old: set[int] = set()
    for _, i, j in cands:
        if i not in pairs and j not in used_old:
            pairs[i] = old[j]
            used_old.add(j)
    return pairs


def _exp_score(n: Experience, o: Experience) -> float | None:
    company = _similar(n.company, o.company)
    if company < 85:
        return None
    title = _similar(n.title, o.title)
    dates = _same_date(n.start, o.start) + _same_date(n.end, o.end)
    if title < 70 and dates < 2:  # same employer, different role
        return None
    return company + title + 60 * dates


def _proj_score(n: Project, o: Project) -> float | None:
    s = _similar(n.name, o.name)
    return s if s >= 85 else None


def _edu_score(n: Education, o: Education) -> float | None:
    school = _similar(n.school, o.school)
    if school < 85:
        return None
    degree = _similar(f"{n.degree} {n.field}", f"{o.degree} {o.field}")
    dates = _same_date(n.start, o.start) + _same_date(n.end, o.end)
    return None if degree < 70 and dates < 2 else school + degree + 60 * dates


def _cert_score(n: Certification, o: Certification) -> float | None:
    s = _similar(n.name, o.name)
    return s if s >= 90 else None


def _merge_items(new: list[ProfileItem], old: list[ProfileItem], ids: _Ids,
                 prefix: str | None) -> None:
    """Carry ids + user fields (tags/context/strength) from matching `old` items, one-to-one.
    Unmatched items get fresh ids (`<prefix>-bN` for bullets)."""
    remaining = list(old)
    matched: dict[int, ProfileItem] = {}
    for item in new:
        best = max(remaining, key=lambda o: _similar(o.text, item.text), default=None)
        if best is not None and _similar(best.text, item.text) < 88:
            best = next((o for o in remaining if o.id == item.id
                         and _similar(o.text, item.text) >= 70), None)
        if best is not None:
            remaining.remove(best)
            matched[id(item)] = best
    for item in new:
        prev = matched.get(id(item))
        if prev is not None and prev.id not in ids.taken:
            item.id = ids.adopt(prev.id)
            item.tags = prev.tags + [t for t in item.tags if t not in prev.tags]
            item.context = prev.context or item.context
            item.strength = prev.strength
        elif prefix is not None:
            item.id = ids.next_bullet(prefix)
        else:
            item.id = ids.fresh(item.id)


def merge_profiles(new: Profile, old: Profile) -> Profile:
    """Carry ids and user-curated fields from `old` onto `new`.

    Parents (experience/projects/education/certifications) are matched one-to-one on
    company + title + dates (name for projects/certs); an old entry is never reused for two
    new ones. Unmatched entries get fresh ids that never collide with any old id.
    """
    new = new.model_copy(deep=True)
    ids = _Ids(_all_ids(old))

    for group, olds, score in ((new.experience, old.experience, _exp_score),
                               (new.projects, old.projects, _proj_score)):
        pairs = _pair(group, olds, score)
        for i, entry in enumerate(group):
            prev = pairs.get(i)
            if prev is not None and prev.id not in ids.taken:
                entry.id = ids.adopt(prev.id)
                _merge_items(entry.bullets, prev.bullets, ids, entry.id)
                continue
            base = entry.id
            if not ids.free(base):
                label = entry.title if isinstance(entry, Experience) else entry.name
                base = f"{base}-{slugify(label, 30)}"
            entry.id = ids.fresh(base)
            for b in entry.bullets:
                b.id = ids.next_bullet(entry.id)

    for group, olds, score in ((new.education, old.education, _edu_score),
                               (new.certifications, old.certifications, _cert_score)):
        pairs = _pair(group, olds, score)
        for i, entry in enumerate(group):
            prev = pairs.get(i)
            entry.id = (ids.adopt(prev.id) if prev is not None and prev.id not in ids.taken
                        else ids.fresh(entry.id))

    _merge_items(new.awards, old.awards, ids, None)
    _merge_items(new.extra, old.extra, ids, None)
    ensure_unique_ids(new)
    return new


def _all_ids(p: Profile) -> list[str]:
    ids = [e.id for e in p.experience] + [x.id for x in p.projects]
    ids += [x.id for x in p.education] + [x.id for x in p.certifications]
    for group in (p.experience, p.projects):
        ids += [b.id for e in group for b in e.bullets]
    return ids + [x.id for x in p.awards] + [x.id for x in p.extra]


def ensure_unique_ids(profile: Profile) -> None:
    """Raise ValueError if any id occurs twice anywhere in the profile."""
    seen: set[str] = set()
    dups = sorted({i for i in _all_ids(profile) if i in seen or seen.add(i)})
    if dups:
        raise ValueError(f"duplicate ids in profile: {', '.join(dups)}")


# --------------------------------------------------------------------------- yaml io


def dump_profile(profile: Profile) -> str:
    return yaml.safe_dump(profile.model_dump(mode="json"), sort_keys=False, allow_unicode=True,
                          width=100)


def _write_atomic(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.name + ".tmp")
    tmp.write_text(text, encoding="utf-8")
    os.replace(tmp, path)


def save_profile(profile: Profile, path: Path) -> None:
    ensure_unique_ids(profile)
    _write_atomic(path, dump_profile(profile))


def read_profile(path: Path) -> Profile:
    data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    profile = Profile.model_validate(data)
    ensure_unique_ids(profile)
    return profile


def load_profile(paths: Paths) -> Profile:
    path = profile_path(paths)
    if not path.exists():
        raise FileNotFoundError(
            f"{path} not found: put your resume files in {paths.resources / 'resume'}, run the "
            "profile ingest and accept the proposal")
    return read_profile(path)


def profile_diff(old: Profile | None, new: Profile, *, fromfile: str = "profile.yaml (current)",
                 tofile: str = "profile.yaml (new)") -> str:
    a = dump_profile(old).splitlines(keepends=True) if old else []
    b = dump_profile(new).splitlines(keepends=True)
    return "".join(difflib.unified_diff(a, b, fromfile=fromfile, tofile=tofile, n=2))


# --------------------------------------------------------------------------- proposal workflow


def proposed_flags_path(paths: Paths) -> Path:
    return paths.data / "profile.proposed.flags.json"


class BlockingFlagsError(ValueError):
    """The proposal has blocking flags; pass allow_blocking=True to accept anyway."""

    def __init__(self, flags: list[VerifierFlag]):
        self.flags = flags
        super().__init__(f"proposal has {len(flags)} blocking flag(s): "
                         + "; ".join(f"{f.where}: {f.text}" for f in flags[:5]))


def _sha(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def read_proposal_flags(paths: Paths) -> list[VerifierFlag] | None:
    """Flags stored with the current proposal; None if there is no (valid) record for it."""
    fp, prop = proposed_flags_path(paths), proposed_profile_path(paths)
    if not fp.exists() or not prop.exists():
        return None
    record = json.loads(fp.read_text(encoding="utf-8"))
    if record.get("proposal_sha256") != _sha(prop.read_text(encoding="utf-8")):
        return None  # the proposal was edited after it was checked
    return [VerifierFlag.model_validate(f) for f in record.get("flags", [])]


class ProposalChanged(ValueError):
    """The proposal on disk isn't the one that was reviewed."""


def proposal_digest(text: str) -> str:
    """Digest of the proposal as reviewed in the UI (sha256 of its bytes, 24 hex chars)."""
    import hashlib

    return hashlib.sha256(text.encode("utf-8")).hexdigest()[:24]


def accept_proposed(paths: Paths, *, allow_blocking: bool = False,
                    expected_digest: str | None = None) -> Profile:
    """Promote data/profile.proposed.yaml to data/profile.yaml.

    Refuses (BlockingFlagsError) when the proposal has blocking flags, or has no valid check
    record (missing, or the proposal was edited afterwards), unless allow_blocking=True. With
    `expected_digest` (the reviewed proposal), a different proposal is refused
    (ProposalChanged). Everything happens under the proposal lock on ONE snapshot of the
    proposal bytes, so flags are validated against exactly what gets promoted. The previous
    profile.yaml is kept as profile.yaml.bak.
    """
    from recrute.tailor.answers import _file_lock

    prop = proposed_profile_path(paths)
    if not prop.exists():
        raise FileNotFoundError(f"no proposal at {prop}")
    with _file_lock(prop):
        text = prop.read_text(encoding="utf-8")
        if expected_digest is not None and proposal_digest(text) != expected_digest:
            raise ProposalChanged("the proposal changed since it was reviewed")
        profile = Profile.model_validate(yaml.safe_load(text) or {})
        ensure_unique_ids(profile)
        flags = None
        fp = proposed_flags_path(paths)
        if fp.exists():
            record = json.loads(fp.read_text(encoding="utf-8"))
            if record.get("proposal_sha256") == _sha(text):
                flags = [VerifierFlag.model_validate(f) for f in record.get("flags", [])]
        if flags is None:
            flags = [VerifierFlag(where="profile.proposed", text="", severity="block",
                                  reason="no check record for this proposal (missing, or the "
                                         "proposal was edited after ingest)")]
        blocking = [f for f in flags if f.severity == "block"]
        if blocking and not allow_blocking:
            raise BlockingFlagsError(blocking)
        target = profile_path(paths)
        if target.exists():
            shutil.copyfile(target, target.with_name(target.name + ".bak"))
        _write_atomic(target, dump_profile(profile))
        prop.unlink()
        fp.unlink(missing_ok=True)
    return profile


# --------------------------------------------------------------------------- entry point


@dataclass
class IngestResult:
    profile: Profile
    flags: list[VerifierFlag] = field(default_factory=list)
    diff: str = ""  # unified diff vs the current profile.yaml ("" when there is none)
    written_to: Path | None = None
    sources: list[str] = field(default_factory=list)
    removed_ids: list[str] = field(default_factory=list)
    accepted: bool = False  # True when promoted to profile.yaml (apply=True, no blocking flags)


def ingest_resume(paths: Paths, router: Completer, *, apply: bool = False) -> IngestResult:
    """Structure resources/resume/* into a proposal for review.

    Always writes data/profile.proposed.yaml plus data/profile.proposed.flags.json (flags,
    diff, removed ids) and leaves profile.yaml alone. Promote with `accept_proposed()`.
    apply=True additionally accepts the proposal right away, but only if it has no blocking
    flags (otherwise it stays a proposal and `accepted` is False).
    """
    sources = read_sources(paths.resources / "resume")
    if not sources:
        raise FileNotFoundError(f"no resume files (.md/.txt/.docx/.pdf) in "
                                f"{paths.resources / 'resume'}")
    source_text = "\n\n".join(f"=== {name} ===\n{text}" for name, text in sources)
    raw = router.complete("extract", EXTRACT_PROMPT.format(source=source_text),
                          schema=EXTRACT_SCHEMA, system=EXTRACT_SYSTEM)
    profile = to_profile(parse_llm(_Extracted, raw, "profile ingest"))
    ensure_unique_ids(profile)

    target = profile_path(paths)
    old = read_profile(target) if target.exists() else None
    removed: list[str] = []
    trusted: set[str] = set()
    if old is not None:
        profile = merge_profiles(profile, old)
        removed = sorted(set(_all_ids(old)) - set(_all_ids(profile)))
        trusted = {i.context for i in old.all_items().values() if i.context}

    flags = check_against_source(profile, source_text, trusted=trusted)
    flags += [VerifierFlag(where=f"profile:{rid}", text=rid, severity="warn",
                           reason="present in the current profile.yaml but not in the new "
                                  "ingest (removed or reworded in the source?)")
              for rid in removed]
    diff = profile_diff(old, profile) if old is not None else ""

    from recrute.tailor.answers import _file_lock

    prop = proposed_profile_path(paths)
    text = dump_profile(profile)
    record = {"proposal_sha256": _sha(text), "created_at": datetime.now(UTC).isoformat(),
              "sources": [n for n, _ in sources], "removed_ids": removed, "diff": diff,
              "flags": [f.model_dump(mode="json") for f in flags]}
    with _file_lock(prop):  # proposal + its check record are published together
        _write_atomic(prop, text)
        _write_atomic(proposed_flags_path(paths),
                      json.dumps(record, indent=2, ensure_ascii=False))

    result = IngestResult(profile=profile, flags=flags, diff=diff, written_to=prop,
                          sources=[n for n, _ in sources], removed_ids=removed)
    if apply and not any(f.severity == "block" for f in flags):
        accept_proposed(paths)
        result.written_to, result.accepted = target, True
    return result
